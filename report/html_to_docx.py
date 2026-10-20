"""Konversi HTML laporan (hasil build_report.build_html) menjadi DOCX yang bisa diedit di Word.

Hanya mendukung subset HTML yang dipakai laporan: h1/h2/h3, p, ol/ul/li, pre, table, figure/img,
figcaption, div.eq/.flow/.tbl/.toc, section.cover, serta inline b/i/sub/sup/code/a/br/span.
Heading memakai style bawaan Word (Heading 1-3) sehingga daftar isi otomatis (field TOC) bekerja.
"""
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x5C, 0x46)
TINT = "F1F5F3"
VOID = {"img", "br", "meta", "hr"}
WS = r"[ \t\r\n\f\v]+"   # bukan \s: \s ikut mencocokkan NBSP (U+00A0) yang sengaja dipakai


# ------------------------------------------------------------------ parser -> pohon sederhana
class Node:
    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs or {}), parent, []

    @property
    def cls(self):
        return set((self.attrs.get("class") or "").split())

    def text(self):
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)

    def find(self, pred):
        for c in self.children:
            if isinstance(c, Node):
                if pred(c):
                    return c
                r = c.find(pred)
                if r:
                    return r
        return None

    def find_all(self, pred):
        out = []
        for c in self.children:
            if isinstance(c, Node):
                if pred(c):
                    out.append(c)
                out.extend(c.find_all(pred))
        return out


class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.cur)
        self.cur.children.append(node)
        if tag not in VOID:
            self.cur = node

    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag:
            n = n.parent
        if n is not self.root:
            self.cur = n.parent

    def handle_data(self, data):
        self.cur.children.append(data)


def parse(html):
    tb = TreeBuilder()
    tb.feed(html)
    body = tb.root.find(lambda n: n.tag == "body")
    return body or tb.root


# ------------------------------------------------------------------ utilitas OOXML
def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths_pt, color) atau None untuk 'nil'."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
            el.set(qn("w:space"), "0")
        borders.append(el)


def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        tcPr.append(w)
    w.set(qn("w:w"), str(int(dxa)))
    w.set(qn("w:type"), "dxa")


def set_table_layout(table, widths_dxa):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths_dxa))))
    tblW.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_dxa):
        gc.set(qn("w:w"), str(int(w)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_dxa):
            set_cell_width(cell, w)


def shade_paragraph(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_field(paragraph, instr, placeholder=""):
    """Sisipkan field Word (PAGE, TOC, ...)."""
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    run._r.append(fld)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run._r.append(it)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    if placeholder:
        paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_hyperlink(paragraph, url, text, fmt):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5C46")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if fmt.get("bold"):
        rPr.append(OxmlElement("w:b"))
    if fmt.get("size"):
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(fmt["size"] * 2)))
        rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# ------------------------------------------------------------------ konverter
class Converter:
    def __init__(self, base_dir):
        self.base = Path(base_dir)
        self.doc = Document()
        self._setup()
        self.list_num_id = None

    # ---------- gaya dokumen
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.4)
        sec.left_margin, sec.right_margin = Cm(3), Cm(2.3)
        sec.different_first_page_header_footer = True
        self.text_width_cm = 21 - 3 - 2.3

        st = d.styles
        normal = st["Normal"]
        normal.font.name = BODY_FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        normal.font.size = Pt(12)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)

        for name, size, italic in (("Heading 1", 14, False), ("Heading 2", 13, False), ("Heading 3", 12, True)):
            h = st[name]
            h.font.name = BODY_FONT
            rpr = h.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), BODY_FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = italic
            h.font.color.rgb = RGBColor(0x1B, 0x1B, 0x1B)
            h.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
            h.paragraph_format.space_after = Pt(6 if name != "Heading 1" else 18)
            h.paragraph_format.line_spacing = 1.15
            h.paragraph_format.keep_with_next = True
        h1 = st["Heading 1"]
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.page_break_before = True

        # judul tanpa masuk daftar isi (DAFTAR ISI itu sendiri)
        plain = st.add_style("Judul Halaman", WD_STYLE_TYPE.PARAGRAPH)
        plain.base_style = normal
        plain.font.size = Pt(14)
        plain.font.bold = True
        plain.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        plain.paragraph_format.page_break_before = True
        plain.paragraph_format.space_after = Pt(18)

        cap = st.add_style("Keterangan", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = normal
        cap.font.size = Pt(10.5)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.space_after = Pt(10)

        code = st.add_style("Kode", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        code.font.name = MONO_FONT
        code.font.size = Pt(9)
        code.paragraph_format.line_spacing = 1.1
        code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        code.paragraph_format.space_after = Pt(8)

        cell = st.add_style("Isi Tabel", WD_STYLE_TYPE.PARAGRAPH)
        cell.base_style = normal
        cell.font.size = Pt(9.5)
        cell.paragraph_format.line_spacing = 1.0
        cell.paragraph_format.space_after = Pt(0)
        cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        for name in ("List Bullet", "List Number"):
            ls = st[name]
            ls.font.name = BODY_FONT
            ls.paragraph_format.line_spacing = 1.5
            ls.paragraph_format.space_after = Pt(3)

        # gaya entri daftar isi bawaan Word ("toc 1", "toc 2") dibuat rapat
        styles_el = d.styles.element
        for level, bold, indent in ((1, True, 0), (2, False, 360)):
            el = OxmlElement("w:style")
            el.set(qn("w:type"), "paragraph")
            el.set(qn("w:styleId"), f"TOC{level}")
            for tag, val in (("w:name", f"toc {level}"), ("w:basedOn", "Normal"), ("w:next", "Normal"),
                             ("w:uiPriority", "39")):
                child = OxmlElement(tag)
                child.set(qn("w:val"), val)
                el.append(child)
            el.append(OxmlElement("w:unhideWhenUsed"))
            ppr = OxmlElement("w:pPr")
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:before"), "60" if level == 1 else "0")
            sp.set(qn("w:after"), "40")
            sp.set(qn("w:line"), "264")
            sp.set(qn("w:lineRule"), "auto")
            ppr.append(sp)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(indent))
            ppr.append(ind)
            el.append(ppr)
            if bold:
                rpr = OxmlElement("w:rPr")
                rpr.append(OxmlElement("w:b"))
                el.append(rpr)
            styles_el.append(el)

        # nomor halaman di footer (halaman sampul tanpa nomor)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE", "1")
        for r in fp.runs:
            r.font.size = Pt(10.5)

    # ---------- inline
    def inline(self, p, node, fmt=None, size=None):
        fmt = dict(fmt or {})
        if size:
            fmt["size"] = size
        for c in node.children:
            if isinstance(c, str):
                txt = re.sub(WS, " ", c)
                if not txt:
                    continue
                r = p.add_run(txt)
                self._fmt(r, fmt)
                continue
            tag = c.tag
            f = dict(fmt)
            if tag in ("b", "strong"):
                f["bold"] = True
            elif tag in ("i", "em"):
                f["italic"] = True
            elif tag == "sub":
                f["sub"] = True
            elif tag == "sup":
                f["sup"] = True
            elif tag == "code":
                f["mono"] = True
            elif tag == "br":
                p.add_run().add_break(WD_BREAK.LINE)
                continue
            elif tag == "a":
                add_hyperlink(p, c.attrs.get("href", ""), re.sub(WS, " ", c.text()), f)
                continue
            elif tag == "span" and "dots" in c.cls:
                continue
            self.inline(p, c, f)

    @staticmethod
    def _fmt(r, f):
        if f.get("bold"):
            r.bold = True
        if f.get("italic"):
            r.italic = True
        if f.get("sub"):
            r.font.subscript = True
        if f.get("sup"):
            r.font.superscript = True
        if f.get("mono"):
            r.font.name = MONO_FONT
            r.font.size = Pt(10)
        if f.get("size"):
            r.font.size = Pt(f["size"])
        if f.get("color"):
            r.font.color.rgb = f["color"]

    @staticmethod
    def _trim(p):
        """Buang spasi di awal/akhir paragraf (sisa whitespace HTML)."""
        runs = [r for r in p.runs if r.text]
        if runs:
            runs[0].text = runs[0].text.lstrip()
            # hyperlink bukan bagian p.runs: jangan potong spasi sebelum link di akhir paragraf
            if runs[-1]._r.getnext() is None:
                runs[-1].text = runs[-1].text.rstrip()

    def para(self, node, style=None, align=None, size=None, fmt=None):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        elif style is None:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.inline(p, node, fmt, size)
        self._trim(p)
        return p

    # ---------- daftar bernomor yang mulai ulang dari 1
    def _new_number_list(self):
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style = self.doc.styles["List Number"]
        base_num_id = int(style.element.pPr.numPr.numId.val)
        abstract_id = None
        for num in numbering.findall(qn("w:num")):
            if int(num.get(qn("w:numId"))) == base_num_id:
                abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), abstract_id)
        num.append(an)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        numbering.append(num)
        return new_id

    def list_block(self, node, numbered):
        num_id = self._new_number_list() if numbered else None
        for li in node.children:
            if not isinstance(li, Node) or li.tag != "li":
                continue
            p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if num_id:
                pPr = p._p.get_or_add_pPr()
                numPr = pPr.get_or_add_numPr()
                numPr.get_or_add_ilvl().val = 0
                numPr.get_or_add_numId().val = num_id
            # konten li: inline + kemungkinan <pre> di dalamnya
            inline_part = Node("span")
            for c in li.children:
                if isinstance(c, Node) and c.tag == "pre":
                    continue
                inline_part.children.append(c)
            self.inline(p, inline_part)
            self._trim(p)
            for c in li.children:
                if isinstance(c, Node) and c.tag == "pre":
                    self.pre(c)

    def pre(self, node):
        lines = node.text().strip("\n").split("\n")
        p = self.doc.add_paragraph(style="Kode")
        shade_paragraph(p, TINT)
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = MONO_FONT
            if i < len(lines) - 1:
                r.add_break(WD_BREAK.LINE)

    # ---------- tabel
    def table(self, node, caption=None, small=False):
        rows = []
        for tr in node.find_all(lambda n: n.tag == "tr"):
            cells = [c for c in tr.children if isinstance(c, Node) and c.tag in ("th", "td")]
            rows.append(cells)
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        numeric = "num" in node.cls
        if caption is not None:
            cp = self.doc.add_paragraph(style="Keterangan")
            cp.paragraph_format.space_after = Pt(3)
            cp.paragraph_format.keep_with_next = True
            self.inline(cp, caption)
            self._trim(cp)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = self._col_widths(rows, ncol, small)
        size = 9 if small else 9.5
        for i, r in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                p.style = self.doc.styles["Isi Tabel"]
                if j < len(r):
                    self.inline(p, r[j], {"bold": r[j].tag == "th"}, size=size)
                    self._trim(p)
                if numeric and j > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                top = (12, "222222") if i == 0 else None
                bottom = (8, "222222") if i == 0 else ((12, "222222") if i == len(rows) - 1 else (4, "D9D9D9"))
                set_cell_borders(cell, top=top, bottom=bottom, left=None, right=None)
        set_table_layout(t, widths)
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i < len(t.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        # baris judul berulang di halaman berikutnya
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 1.0

    def _col_widths(self, rows, ncol, small):
        """Lebar kolom (DXA): minimal muat kata terpanjang, sisa ruang dibagi
        proporsional dengan panjang isi sel (dibatasi) agar tidak ada kata terpotong."""
        char = 92 if small else 98          # ±lebar rata-rata 1 karakter Times 9-9,5 pt
        pad = 230                           # margin sel kiri+kanan
        total = self.text_width_cm * 567
        minw, want = [pad] * ncol, [pad] * ncol
        for r in rows:
            for j, c in enumerate(r):
                txt = re.sub(WS, " ", c.text()).strip()
                k = char * (1.12 if c.tag == "th" else 1.0)
                if c.find(lambda n: n.tag == "code"):
                    k *= 1.25
                # huruf kapital Times ±1,45x lebih lebar dari huruf kecil
                units = lambda w: sum(1.45 if ch.isupper() else 1.0 for ch in w)  # noqa: E731
                longest = max((units(w) for w in txt.split(" ")), default=0)
                minw[j] = max(minw[j], longest * k + pad)
                want[j] = max(want[j], min(units(txt), 42) * k + pad)
        if sum(minw) >= total:
            return [w * total / sum(minw) for w in minw]
        if sum(want) <= total:
            return [w * total / sum(want) for w in want]
        extra = total - sum(minw)
        flex = [want[j] - minw[j] for j in range(ncol)]
        return [minw[j] + extra * flex[j] / sum(flex) for j in range(ncol)]

    def flow(self, node):
        boxes = [c for c in node.children if isinstance(c, Node) and c.tag == "div"]
        cols = 3
        nrow = (len(boxes) + cols - 1) // cols
        t = self.doc.add_table(rows=nrow, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, b in enumerate(boxes):
            cell = t.cell(k // cols, k % cols)
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Isi Tabel"]
            self.inline(p, b, size=10)
            self._trim(p)
            set_cell_shading(cell, TINT)
            set_cell_borders(cell, top=(4, "FFFFFF"), bottom=(4, "FFFFFF"), left=(4, "FFFFFF"), right=(4, "FFFFFF"))
        set_table_layout(t, [self.text_width_cm * 567 / cols] * cols)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---------- gambar
    def figure(self, node):
        img = node.find(lambda n: n.tag == "img")
        cap = node.find(lambda n: n.tag == "figcaption")
        m = re.search(r"width:\s*([\d.]+)%", img.attrs.get("style", ""))
        frac = float(m.group(1)) / 100 if m else 1.0
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str((self.base / img.attrs["src"]).resolve()), width=Cm(self.text_width_cm * frac))
        if cap is not None:
            self.para(cap, style="Keterangan")

    # ---------- sampul
    def cover(self, node):
        def by_cls(c):
            return node.find(lambda n: c in n.cls)

        def centered(text_node, size, bold=False, italic=False, color=None, space_after=6, caps=False):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(space_after)
            self.inline(p, text_node, {"bold": bold, "italic": italic, "color": color}, size=size)
            self._trim(p)
            if caps:
                for r in p.runs:
                    r.text = r.text.upper()
            return p

        gap = lambda pt: self.doc.add_paragraph().paragraph_format.__setattr__("space_after", Pt(pt))  # noqa: E731
        centered(by_cls("cover-kicker"), 11, bold=True, color=ACCENT, space_after=36, caps=True)
        centered(by_cls("cover-title"), 20, bold=True, space_after=14)
        centered(by_cls("cover-sub"), 13, italic=True, space_after=60)
        box = by_cls("cover-box")
        labels = [n for n in box.children if isinstance(n, Node) and "cover-label" in n.cls]
        team = box.find(lambda n: n.tag == "table")
        centered(labels[0], 10, bold=True, color=RGBColor(0x55, 0x55, 0x55), space_after=4, caps=True)
        for tr in team.find_all(lambda n: n.tag == "tr"):
            tds = [c for c in tr.children if isinstance(c, Node)]
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            p.add_run(tds[0].text().strip()).font.size = Pt(12.5)
            r = p.add_run("    " + tds[1].text().strip())
            r.font.size = Pt(12.5)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        gap(10)
        centered(labels[1], 10, bold=True, color=RGBColor(0x55, 0x55, 0x55), space_after=4, caps=True)
        lect = [n for n in box.children if isinstance(n, Node) and n.tag == "div" and not n.cls]
        centered(lect[-1], 12.5, space_after=170)
        centered(by_cls("cover-foot"), 13, bold=True, space_after=0)

    # ---------- blok utama
    def block(self, node):
        if isinstance(node, str):
            if node.strip():
                p = self.doc.add_paragraph()
                p.add_run(re.sub(WS, " ", node.strip()))
            return
        tag, cls = node.tag, node.cls
        if tag == "section" and "cover" in cls:
            self.cover(node)
        elif tag == "h1":
            if "plain" in cls and node.attrs.get("id") == "toc":
                p = self.doc.add_paragraph(style="Judul Halaman")
                p.add_run(node.text().strip())
            else:
                p = self.doc.add_paragraph(style="Heading 1")
                span = node.find(lambda n: n.tag == "span")
                if span is not None and "plain" not in cls:
                    # satu baris ("BAB 1 PENDAHULUAN") agar entri daftar isi tidak terpecah
                    rest = "".join(c for c in node.children if isinstance(c, str)).strip()
                    p.add_run(f"{span.text().strip()}  {rest}")
                else:
                    p.add_run(node.text().strip())
        elif tag == "h2":
            self.para(node, style="Heading 2", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif tag == "h3":
            self.para(node, style="Heading 3", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif tag == "p":
            if "gh" in cls:
                p = self.para(node, align=WD_ALIGN_PARAGRAPH.LEFT, fmt={"bold": True}, size=11)
                shade_paragraph(p, TINT)
            elif "kw" in cls:
                self.para(node, size=11)
            else:
                self.para(node)
        elif tag == "ol" and "refs" in cls:
            items = [c for c in node.children if isinstance(c, Node) and c.tag == "li"]
            for k, li in enumerate(items, start=1):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf = p.paragraph_format
                pf.left_indent, pf.first_line_indent = Cm(1.0), Cm(-1.0)
                pf.line_spacing, pf.space_after = 1.15, Pt(5)
                r = p.add_run(f"[{k}]\t")
                r.font.size = Pt(11)
                self.inline(p, li, size=11)
                self._trim(p)
                p.runs[0].text = f"[{k}]\t"
        elif tag in ("ol", "ul"):
            self.list_block(node, numbered=(tag == "ol"))
        elif tag == "pre":
            self.pre(node)
        elif tag == "figure":
            self.figure(node)
        elif tag == "div" and "eq" in cls:
            p = self.para(node, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(8)
        elif tag == "div" and "flow" in cls:
            self.flow(node)
        elif tag == "div" and "tbl" in cls:
            cap = node.find(lambda n: "tcap" in n.cls)
            tbl = node.find(lambda n: n.tag == "table")
            self.table(tbl, caption=cap, small="small" in tbl.cls)
        elif tag == "div" and "toc" in cls:
            p = self.doc.add_paragraph()
            add_field(p, 'TOC \\o "1-2" \\h \\z \\u',
                      "Klik kanan → Update Field untuk menampilkan daftar isi.")
        else:  # kontainer lain: telusuri anaknya
            for c in node.children:
                self.block(c)

    def convert(self, html, out_path):
        body = parse(html)
        for c in body.children:
            self.block(c)
        # buang paragraf kosong pertama bawaan template (jika ada)
        first = self.doc.paragraphs[0]
        if not first.text and not first.runs:
            first._p.getparent().remove(first._p)
        self.doc.core_properties.title = "Optimasi Rute Distribusi MBG Surabaya"
        self.doc.core_properties.subject = "Final Project Konsep Kecerdasan Artifisial"
        self.doc.save(str(out_path))
        return out_path


def html_to_docx(html, out_path, base_dir):
    return Converter(base_dir).convert(html, out_path)
